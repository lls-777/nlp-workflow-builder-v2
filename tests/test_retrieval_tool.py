from __future__ import annotations

from pathlib import Path
from typing import Any, List

import pytest
from docx import Document as DocxDocument

from tools import documents
from tools.documents import DocumentReadError, read_document_file
from tools.retrieval import chunk_documents, retrieve_relevant_chunks


class FakePdfPage:
    def __init__(self, text: str) -> None:
        self.text = text

    def extract_text(self) -> str:
        return self.text


class FakePdfReader:
    def __init__(self, texts: List[str], *, encrypted: bool = False) -> None:
        self.pages = [FakePdfPage(text) for text in texts]
        self.is_encrypted = encrypted


def test_read_pdf_with_text_layer(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    pdf_path = tmp_path / "guide.pdf"
    pdf_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        documents,
        "PdfReader",
        lambda path: FakePdfReader(["First PDF page.", "Second PDF page."]),
    )

    result = read_document_file(pdf_path)

    assert result["content"] == "First PDF page.\n\nSecond PDF page."


def test_read_pdf_rejects_encrypted_file(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    pdf_path = tmp_path / "encrypted.pdf"
    pdf_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        documents,
        "PdfReader",
        lambda path: FakePdfReader(["Hidden content"], encrypted=True),
    )

    with pytest.raises(DocumentReadError, match="Encrypted PDF"):
        read_document_file(pdf_path)


def test_read_pdf_rejects_damaged_file(tmp_path: Path) -> None:
    pdf_path = tmp_path / "damaged.pdf"
    pdf_path.write_bytes(b"not a valid PDF")

    with pytest.raises(DocumentReadError, match="damaged or cannot be read"):
        read_document_file(pdf_path)


def test_read_pdf_rejects_file_without_text_layer(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    pdf_path = tmp_path / "scan.pdf"
    pdf_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        documents,
        "PdfReader",
        lambda path: FakePdfReader(["", "  "]),
    )

    with pytest.raises(DocumentReadError, match="OCR"):
        read_document_file(pdf_path)


def test_read_docx_paragraphs_and_table_cells(tmp_path: Path) -> None:
    docx_path = tmp_path / "faq.docx"
    document = DocxDocument()
    document.add_paragraph("Membership FAQ")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Refund"
    table.cell(0, 1).text = "Available within 14 days"
    document.save(docx_path)

    result = read_document_file(docx_path)

    assert result["content"] == (
        "Membership FAQ\nRefund\nAvailable within 14 days"
    )


def test_chunk_documents_uses_stable_ids_and_overlap() -> None:
    chunks = chunk_documents(
        [
            {
                "name": "faq.md",
                "path": "/internal/faq.md",
                "content": "abcdefghij",
            }
        ],
        chunk_size=6,
        overlap=2,
    )

    assert chunks == [
        {
            "chunk_id": "faq.md::chunk_0001",
            "source_name": "faq.md",
            "source_path": "/internal/faq.md",
            "content": "abcdef",
            "start_char": 0,
            "end_char": 6,
        },
        {
            "chunk_id": "faq.md::chunk_0002",
            "source_name": "faq.md",
            "source_path": "/internal/faq.md",
            "content": "efghij",
            "start_char": 4,
            "end_char": 10,
        },
    ]


@pytest.mark.parametrize(
    ("chunks", "query"),
    [
        ([], "refund"),
        ([{"content": ""}], "refund"),
        ([{"content": "refund policy"}], ""),
        ([{"content": "a"}], "a"),
    ],
)
def test_retrieve_chunks_handles_empty_inputs_without_traceback(
    chunks: List[dict], query: str
) -> None:
    assert retrieve_relevant_chunks(chunks, query=query, top_k=5) == []


def test_retrieve_chunks_returns_relevant_results_in_stable_order() -> None:
    chunks = [
        {
            "chunk_id": "faq.md::chunk_0001",
            "source_name": "faq.md",
            "source_path": "/internal/faq.md",
            "content": "Refund requests are accepted within 14 days.",
        },
        {
            "chunk_id": "faq.md::chunk_0002",
            "source_name": "faq.md",
            "source_path": "/internal/faq.md",
            "content": "Premium members receive priority email support.",
        },
    ]

    results = retrieve_relevant_chunks(chunks, query="refund request", top_k=1)

    assert len(results) == 1
    assert results[0]["chunk_id"] == "faq.md::chunk_0001"
    assert results[0]["score"] > 0
