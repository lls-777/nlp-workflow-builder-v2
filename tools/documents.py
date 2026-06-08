from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Sequence, Union

from docx import Document as DocxDocument
from pypdf import PdfReader

from config import MAX_EXTRACTED_TEXT_CHARS, MAX_PDF_PAGES
from core.workflow_schema import ExecutionContext


TEXT_EXTENSIONS = frozenset({".txt", ".md"})
PDF_EXTENSION = ".pdf"
DOCX_EXTENSION = ".docx"


class DocumentReadError(ValueError):
    """Raised when a document cannot be read in the current stage."""


def read_document_file(file_path: Union[str, Path]) -> Dict[str, str]:
    path = Path(file_path).resolve()
    if not path.is_file():
        raise DocumentReadError(f"Document file does not exist: {path}")

    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        content = _read_text_document(path)
    elif extension == PDF_EXTENSION:
        content = _read_pdf_document(path)
    elif extension == DOCX_EXTENSION:
        content = _read_docx_document(path)
    else:
        raise DocumentReadError(
            f"Unsupported document format: {extension or '<none>'}. "
            "Supported formats are TXT, MD, PDF, and DOCX."
        )

    _validate_extracted_text(path, content)
    return {
        "name": path.name,
        "path": str(path),
        "content": content,
    }


def _read_text_document(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentReadError(
            f"TXT and MD documents must use UTF-8 encoding: {path.name}"
        ) from exc
    except OSError as exc:
        raise DocumentReadError(f"Unable to read document: {path.name}") from exc


def _read_pdf_document(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentReadError(
            f"PDF file is damaged or cannot be read: {path.name}"
        ) from exc

    if reader.is_encrypted:
        raise DocumentReadError(
            f"Encrypted PDF files are not supported: {path.name}"
        )
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentReadError(
            f"PDF file exceeds the {MAX_PDF_PAGES} page limit: {path.name}"
        )

    page_texts = []
    try:
        for page in reader.pages:
            page_texts.append(page.extract_text() or "")
    except Exception as exc:
        raise DocumentReadError(
            f"Unable to extract text from PDF file: {path.name}"
        ) from exc

    content = "\n\n".join(page_texts).strip()
    if not content:
        raise DocumentReadError(
            f"PDF file has no readable text layer: {path.name}. "
            "Scanned PDFs require OCR, which is not supported in the current MVP."
        )
    return content


def _read_docx_document(path: Path) -> str:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentReadError(
            f"DOCX file is damaged or cannot be read: {path.name}"
        ) from exc

    parts = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text.strip() for cell in row.cells)
    return "\n".join(part for part in parts if part)


def _validate_extracted_text(path: Path, content: str) -> None:
    if not content.strip():
        raise DocumentReadError(f"Document has no readable text: {path.name}")
    if len(content) > MAX_EXTRACTED_TEXT_CHARS:
        raise DocumentReadError(
            "Extracted document text exceeds the "
            f"{MAX_EXTRACTED_TEXT_CHARS} character limit: {path.name}"
        )


def read_documents(file_paths: Sequence[Union[str, Path]]) -> List[Dict[str, str]]:
    if not file_paths:
        raise DocumentReadError("At least one document file is required.")
    return [read_document_file(file_path) for file_path in file_paths]


def read_document(
    inputs: Dict[str, Any],
    params: Dict[str, Any],
    context: ExecutionContext,
) -> Dict[str, Any]:
    files = inputs.get("files")
    if isinstance(files, (str, bytes, Path)) or not isinstance(files, Sequence):
        raise DocumentReadError("read_document requires a list of document files.")
    return {"documents": read_documents(files)}
